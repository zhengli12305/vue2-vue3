import io
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Dict, List, Tuple

import docx
import openpyxl
import pdfplumber

from .schemas import ParseQuestionsResponse, QuizOption, QuizQuestion

OPTION_RE = re.compile(r"^[A-H][\.\、\)]\s*(.+)$")

# 更宽松的题目标题识别 (匹配 "一、单项选择题", "【多选】", "判断题")
TYPE_SECTION_RE = re.compile(r"^(?:[一二三四五六七八九十]、)?\s*([【\[]?(单项选择|多项选择|判断|单选|多选|问答)[】\]]?(?:题)?)\s*$")
# 题目识别 (匹配 "1. 题目", "1、题目", "1)题目")
QUESTION_RE = re.compile(r"^\s*(\d+)[\.\、\)]\s*(.+)$")
# 答案识别 (匹配 "答案: A", "答案：B,C", "正确答案：D")
ANSWER_RE = re.compile(
    r"^(?:(?:答案|正确答案|答)|[^\wA-Za-z0-9]{1,4})[\:\：]\s*([A-Za-z0-9,\s对错√✖]+)$",
    flags=re.IGNORECASE
)
# 集中答案区域识别 (匹配 "参考答案", "答案")
ANSWERS_SECTION_START_RE = re.compile(r"^(?:参考答案|答案)$", flags=re.IGNORECASE)

# 集中答案行识别 (匹配 "1. A", "2. B,C")
CENTRALIZED_ANSWER_LINE_RE = re.compile(r"^\s*(\d+)[\.\、\)]\s*([A-Za-z0-9,\s对错√✖]+)$", flags=re.IGNORECASE)
BATCH_ANSWER_PAIR_RE = re.compile(r"(\d+)\s*[\.\、\)]?\s*([A-H]+|[A-H](?:[,\s，]+[A-H])+)", flags=re.IGNORECASE)
INLINE_QUESTION_SPLIT_RE = re.compile(r"(?<!^)(?<!\d)(\d{1,3}[\.\、\)])")
INLINE_OPTION_SPLIT_RE = re.compile(r"(?<!^)([A-H][\.\、\)])")

def parse_uploaded_file(filename: str, content: bytes) -> ParseQuestionsResponse:
    suffix = Path(filename).suffix.lower()
    if suffix == ".doc":
        content = _convert_doc_to_docx_content(content)
        lines = _extract_from_word(content)
    elif suffix == ".docx":
        lines = _extract_from_word(content)
    elif suffix == ".pdf":
        lines = _extract_from_pdf(content)
    elif suffix in {".xls", ".xlsx"}:
        lines = _extract_from_excel(content)
    else:
        raise ValueError("文件格式不支持，仅支持 Word/PDF/Excel。")

    # 分阶段解析：先解析题目和选项，再解析集中答案区，最后匹配
    parsed_questions_data = _parse_questions_and_options(lines)
    centralized_answers = _parse_centralized_answers(lines)

    if not parsed_questions_data:
        raise ValueError("未识别出有效题目，请检查题干/选项格式。")

    # 尝试将集中答案匹配到题目中
    questions_with_answers = _associate_answers_to_questions(
        parsed_questions_data, centralized_answers
    )

    if not questions_with_answers:
        raise ValueError("未找到任何有效题目和答案匹配，请检查文档。")

    return ParseQuestionsResponse(
        quizTitle=Path(filename).stem, questions=questions_with_answers
    )

def _resolve_soffice_path() -> str | None:
    env_path = os.getenv("SOFFICE_PATH")
    if env_path and Path(env_path).exists():
        return env_path

    which_path = shutil.which("soffice") or shutil.which("soffice.exe")
    if which_path:
        return which_path

    common_paths = [
        Path("C:/Program Files/LibreOffice/program/soffice.exe"),
        Path("C:/Program Files (x86)/LibreOffice/program/soffice.exe"),
    ]
    for candidate in common_paths:
        if candidate.exists():
            return str(candidate)
    return None

def _convert_doc_to_docx_content(content: bytes) -> bytes:
    soffice_path = _resolve_soffice_path()
    if not soffice_path:
        raise ValueError("检测到 .doc 文件，但未找到 LibreOffice。请安装后重试。")

    with tempfile.TemporaryDirectory(prefix="quiz-doc-convert-") as tmpdir:
        tmp_path = Path(tmpdir)
        source_file = tmp_path / "source.doc"
        target_file = tmp_path / "source.docx"
        source_file.write_bytes(content)

        command = [
            soffice_path,
            "--headless",
            "--convert-to",
            "docx",
            "--outdir",
            str(tmp_path),
            str(source_file),
        ]

        try:
            completed = subprocess.run(command, capture_output=True, text=True, timeout=60)
        except subprocess.TimeoutExpired as exc:
            raise ValueError("DOC 转换超时，请重试或改为上传 .docx。") from exc
        except OSError as exc:
            raise ValueError("DOC 转换程序调用失败，请检查 LibreOffice 配置。") from exc

        if completed.returncode != 0 or not target_file.exists():
            raise ValueError("DOC 转换失败，请检查文件是否损坏，或改为上传 .docx。")

        return target_file.read_bytes()

def _extract_from_word(content: bytes) -> List[str]:
    document = docx.Document(io.BytesIO(content))
    lines: List[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        # 兼容“一个段落内包含多行题干/选项/答案”的文档
        for piece in re.split(r"[\r\n]+", text):
            piece = piece.strip()
            if piece:
                lines.append(piece)
    return lines

def _extract_from_pdf(content: bytes) -> List[str]:
    lines: List[str] = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            lines.extend([line.strip() for line in text.splitlines() if line.strip()])
    return lines

def _extract_from_excel(content: bytes) -> List[str]:
    workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    lines: List[str] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            for cell in row:
                if cell is None:
                    continue
                value = str(cell).strip()
                if value:
                    lines.append(value)
    return lines

def _parse_questions_and_options(lines: List[str]) -> List[Dict]:
    questions: List[Dict] = []
    current: Dict | None = None
    current_type = "ONE"
    in_answers_tail = False
    expect_stem_after_type = False
    used_ids: set[str] = set()
    next_auto_id = 1

    for raw_line in lines:
        for line in _expand_compound_line(raw_line.strip()):
            if not line:
                continue

            if ANSWERS_SECTION_START_RE.search(line):
                in_answers_tail = True
                continue

            if _is_batch_answer_line(line):
                in_answers_tail = True
                continue

            if in_answers_tail:
                continue

            type_section_match = TYPE_SECTION_RE.match(line)
            if type_section_match:
                type_label = type_section_match.group(2)
                if type_label in {"多项选择", "多选"}:
                    current_type = "MORE"
                elif type_label == "判断":
                    current_type = "JUDGE"
                elif type_label in {"单项选择", "单选"}:
                    current_type = "ONE"
                expect_stem_after_type = True
                continue

            question_match = QUESTION_RE.match(line)
            if question_match:
                if current:
                    questions.append(current)
                qid, stem = question_match.groups()
                qid = str(qid)
                current = {
                    "id": qid,
                    "type": current_type,
                    "stem": stem,
                    "options": [],
                    "correctAnswerIds": [],
                }
                used_ids.add(qid)
                if qid.isdigit():
                    next_auto_id = max(next_auto_id, int(qid) + 1)
                expect_stem_after_type = False
                continue

            if not current:
                # 兼容：题型标记后紧跟“无题号题干”
                if expect_stem_after_type and not OPTION_RE.match(line) and not ANSWER_RE.match(line):
                    while str(next_auto_id) in used_ids:
                        next_auto_id += 1
                    qid = str(next_auto_id)
                    next_auto_id += 1
                    current = {
                        "id": qid,
                        "type": current_type,
                        "stem": line,
                        "options": [],
                        "correctAnswerIds": [],
                    }
                    used_ids.add(qid)
                    expect_stem_after_type = False
                continue

            option_match = OPTION_RE.match(line)
            if option_match:
                label = option_match.group(0)[0].upper()
                current["options"].append({"id": label, "text": option_match.group(1).strip()})
                continue

            # 兼容“每题后紧跟答案”的格式
            answer_match = ANSWER_RE.match(line)
            if answer_match:
                current["correctAnswerIds"] = _normalize_answer_labels(answer_match.group(1))
                continue

            if current["options"] and not (QUESTION_RE.match(line) or ANSWER_RE.match(line) or TYPE_SECTION_RE.match(line)):
                last_option = current["options"][-1]
                last_option["text"] = f"{last_option['text']} {line}".strip()

    if current:
        questions.append(current)
    return questions

def _parse_centralized_answers(lines: List[str]) -> Dict[str, List[str]]:
    centralized_answers: Dict[str, List[str]] = {}
    in_answers_section = False
    for line in lines:
        line = line.strip()
        if ANSWERS_SECTION_START_RE.match(line):
            in_answers_section = True
            continue

        # 先尝试解析“一行多个答案对”，如：1.A 2.B 3.C
        batch_pairs = BATCH_ANSWER_PAIR_RE.findall(line)
        if (in_answers_section and batch_pairs) or len(batch_pairs) >= 2:
            for qid, answers_str in batch_pairs:
                normalized = _normalize_answer_labels(answers_str)
                if normalized:
                    centralized_answers[str(qid)] = normalized
            continue

        if in_answers_section:
            answer_match = CENTRALIZED_ANSWER_LINE_RE.match(line)
            if answer_match:
                qid, answers_str = answer_match.groups()
                centralized_answers[str(qid)] = _normalize_answer_labels(answers_str)
    return centralized_answers


def _is_batch_answer_line(line: str) -> bool:
    return len(BATCH_ANSWER_PAIR_RE.findall(line)) >= 2


def _split_by_markers(line: str, split_re: re.Pattern[str]) -> List[str]:
    if not line:
        return []
    matches = list(split_re.finditer(line))
    if not matches:
        return [line.strip()]
    parts: List[str] = []
    start = 0
    for match in matches:
        idx = match.start()
        if idx > start:
            chunk = line[start:idx].strip()
            if chunk:
                parts.append(chunk)
        start = idx
    tail = line[start:].strip()
    if tail:
        parts.append(tail)
    return parts


def _expand_compound_line(line: str) -> List[str]:
    if not line:
        return []

    normalized = line.replace("\t", " ").strip()
    if _is_batch_answer_line(normalized):
        return [normalized]

    # 先拆“同一行多个题号”
    question_chunks = _split_by_markers(normalized, INLINE_QUESTION_SPLIT_RE)

    # 再拆“同一行多个选项”
    final_chunks: List[str] = []
    for chunk in question_chunks:
        option_chunks = _split_by_markers(chunk, INLINE_OPTION_SPLIT_RE)
        final_chunks.extend(option_chunks)
    return [item.strip() for item in final_chunks if item.strip()]

def _associate_answers_to_questions(
    parsed_questions_data: List[Dict],
    centralized_answers: Dict[str, List[str]],
) -> List[QuizQuestion]:
    final_questions: List[QuizQuestion] = []
    for question_data in parsed_questions_data:
        # 优先从集中答案中获取答案
        qid = str(question_data["id"])
        correct_answers = centralized_answers.get(qid)

        # 如果集中答案中没有，再尝试从题目行中获取 (兼容老格式)
        if not correct_answers and question_data.get("correctAnswerIds"):
            correct_answers = question_data["correctAnswerIds"]

        if correct_answers:
            question_data["correctAnswerIds"] = correct_answers
            try:
                final_questions.append(_validate_question(question_data))
            except ValueError as e:
                print(f"Warning: Skipping question {qid} due to validation error: {e}")
        else:
            print(f"Warning: Skipping question {qid} as no correct answers were found.")
    return final_questions

def _normalize_answer_labels(raw_answers: str) -> List[str]:
    # 兼容对/错，以及多选答案无逗号
    if raw_answers.lower() in {"对", "√", "正确"}:
        return ["A"]
    if raw_answers.lower() in {"错", "✖", "错误"}:
        return ["B"]

    # 尝试按逗号或空格分割，然后按字符分割（兼容ACD -> A,C,D）
    labels = []
    parts = re.split(r"[,，\s]+", raw_answers)
    for part in parts:
        if part.strip():
            # 如果是单个字母，直接添加
            if len(part.strip()) == 1 and part.strip().isalpha():
                labels.append(part.strip().upper())
            # 如果是连续字母，拆开添加 (例如 "ACD" -> "A", "C", "D")
            elif len(part.strip()) > 1 and part.strip().isalpha():
                labels.extend(list(part.strip().upper()))
            else:
                # 其他情况，例如数字答案，直接添加
                labels.append(part.strip())

    return list(dict.fromkeys(labels))

def _validate_question(question: Dict) -> QuizQuestion:
    option_ids = {item["id"] for item in question["options"]}
    answers = [answer for answer in question["correctAnswerIds"] if answer in option_ids]

    if not answers:
        raise ValueError(f"题目 {question['id']} 未识别到有效标准答案。")

    if question["type"] in {"ONE", "JUDGE"} and len(answers) != 1:
        raise ValueError(f"题目 {question['id']} 为单选/判断，标准答案必须唯一。")

    if question["type"] == "MORE" and len(answers) < 1:
        raise ValueError(f"题目 {question['id']} 为多选，标准答案不能为空。")

    return QuizQuestion(
        id=str(question["id"]),
        type=question["type"],
        stem=question["stem"],
        options=[QuizOption(**item) for item in question["options"]],
        correctAnswerIds=answers,
    )
